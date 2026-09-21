from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    Paragraph,
)

from app.services.dashboard_service import (
    get_monthly_booking_report,
    get_resource_usage,
    get_room_utilization,
    get_upcoming_meetings,
)


def create_excel_report(
    db,
    year: int,
    month: int,
) -> BytesIO:

    monthly_report = get_monthly_booking_report(
        db,
        year,
        month,
    )

    room_utilization = get_room_utilization(
        db,
        year,
        month,
    )

    resource_usage = get_resource_usage(
        db,
        year,
        month,
    )

    upcoming_meetings = get_upcoming_meetings(
        db,
        limit=100,
    )

    workbook = Workbook()

    # Summary sheet
    summary_sheet = workbook.active
    summary_sheet.title = "Monthly Summary"

    summary_sheet.append(
        [f"Monthly Booking Report - {year}-{month:02d}"]
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        size=14,
    )

    summary_sheet.append([])
    summary_sheet.append(
        ["Metric", "Value"]
    )

    summary_rows = [
        ["Total Bookings", monthly_report["total_bookings"]],
        [
            "Confirmed Bookings",
            monthly_report["confirmed_bookings"],
        ],
        [
            "Cancelled Bookings",
            monthly_report["cancelled_bookings"],
        ],
        [
            "Recurring Bookings",
            monthly_report["recurring_bookings"],
        ],
        [
            "Total Booked Hours",
            monthly_report["total_booked_hours"],
        ],
    ]

    for row in summary_rows:
        summary_sheet.append(row)

    # Room utilization sheet
    room_sheet = workbook.create_sheet(
        "Room Utilization"
    )

    room_sheet.append(
        [
            "Room ID",
            "Room Name",
            "Room Code",
            "Total Bookings",
            "Booked Hours",
            "Utilization %",
        ]
    )

    for cell in room_sheet[1]:
        cell.font = Font(bold=True)

    for item in room_utilization:
        room_sheet.append(
            [
                item["room_id"],
                item["room_name"],
                item["room_code"],
                item["total_bookings"],
                item["total_booked_hours"],
                item["utilization_percentage"],
            ]
        )

    # Resource usage sheet
    resource_sheet = workbook.create_sheet(
        "Resource Usage"
    )

    resource_sheet.append(
        [
            "Resource ID",
            "Resource Name",
            "Resource Code",
            "Total Quantity Booked",
            "Booking Count",
        ]
    )

    for cell in resource_sheet[1]:
        cell.font = Font(bold=True)

    for item in resource_usage:
        resource_sheet.append(
            [
                item["resource_id"],
                item["resource_name"],
                item["resource_code"],
                item["total_quantity_booked"],
                item["booking_count"],
            ]
        )

    # Upcoming meetings sheet
    meetings_sheet = workbook.create_sheet(
        "Upcoming Meetings"
    )

    meetings_sheet.append(
        [
            "Booking ID",
            "Title",
            "Room ID",
            "User ID",
            "Start",
            "End",
            "Status",
        ]
    )

    for cell in meetings_sheet[1]:
        cell.font = Font(bold=True)

    for booking in upcoming_meetings:
        meetings_sheet.append(
            [
                booking["booking_id"],
                booking["title"],
                booking["room_id"],
                booking["user_id"],
                booking["start_datetime"].strftime(
                    "%Y-%m-%d %H:%M"
                ),
                booking["end_datetime"].strftime(
                    "%Y-%m-%d %H:%M"
                ),
                booking["status"],
            ]
        )

    # Basic column sizing
    for sheet in workbook.worksheets:
        for column in sheet.columns:
            max_length = 0

            for cell in column:
                value = str(cell.value or "")
                max_length = max(
                    max_length,
                    len(value),
                )

            sheet.column_dimensions[
                column[0].column_letter
            ].width = min(
                max_length + 2,
                40,
            )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def create_pdf_report(
    db,
    year: int,
    month: int,
) -> BytesIO:

    monthly_report = get_monthly_booking_report(
        db,
        year,
        month,
    )

    room_utilization = get_room_utilization(
        db,
        year,
        month,
    )

    resource_usage = get_resource_usage(
        db,
        year,
        month,
    )

    upcoming_meetings = get_upcoming_meetings(
        db,
        limit=100,
    )

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25,
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            f"Smart Meeting Room & Resource Management "
            f"System - Monthly Report",
            styles["Title"],
        )
    )

    elements.append(
        Paragraph(
            f"Reporting Period: {year}-{month:02d}",
            styles["Heading2"],
        )
    )

    elements.append(Spacer(1, 12))

    # Summary
    summary_data = [
        ["Metric", "Value"],
        [
            "Total Bookings",
            str(monthly_report["total_bookings"]),
        ],
        [
            "Confirmed Bookings",
            str(monthly_report["confirmed_bookings"]),
        ],
        [
            "Cancelled Bookings",
            str(monthly_report["cancelled_bookings"]),
        ],
        [
            "Recurring Bookings",
            str(monthly_report["recurring_bookings"]),
        ],
        [
            "Total Booked Hours",
            str(monthly_report["total_booked_hours"]),
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[200, 120],
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.black,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "ALIGN",
                    (1, 1),
                    (1, -1),
                    "RIGHT",
                ),
            ]
        )
    )

    elements.append(
        Paragraph(
            "Monthly Summary",
            styles["Heading2"],
        )
    )

    elements.append(summary_table)
    elements.append(Spacer(1, 18))

    # Room utilization
    room_data = [
        [
            "Room",
            "Code",
            "Bookings",
            "Hours",
            "Utilization %",
        ]
    ]

    for item in room_utilization:
        room_data.append(
            [
                item["room_name"],
                item["room_code"],
                str(item["total_bookings"]),
                str(item["total_booked_hours"]),
                str(item["utilization_percentage"]),
            ]
        )

    elements.append(
        Paragraph(
            "Room Utilization",
            styles["Heading2"],
        )
    )

    room_table = Table(
        room_data,
        repeatRows=1,
    )

    room_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
            ]
        )
    )

    elements.append(room_table)
    elements.append(Spacer(1, 18))

    # Resource usage
    resource_data = [
        [
            "Resource",
            "Code",
            "Quantity Booked",
            "Booking Count",
        ]
    ]

    for item in resource_usage:
        resource_data.append(
            [
                item["resource_name"],
                item["resource_code"],
                str(item["total_quantity_booked"]),
                str(item["booking_count"]),
            ]
        )

    elements.append(
        Paragraph(
            "Resource Usage",
            styles["Heading2"],
        )
    )

    resource_table = Table(
        resource_data,
        repeatRows=1,
    )

    resource_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
            ]
        )
    )

    elements.append(resource_table)
    elements.append(Spacer(1, 18))

    # Upcoming meetings
    meeting_data = [
        [
            "Booking",
            "Title",
            "Room",
            "Start",
            "End",
            "Status",
        ]
    ]

    for booking in upcoming_meetings:
        meeting_data.append(
            [
                str(booking["booking_id"]),
                booking["title"],
                str(booking["room_id"]),
                booking["start_datetime"].strftime(
                    "%Y-%m-%d %H:%M"
                ),
                booking["end_datetime"].strftime(
                    "%Y-%m-%d %H:%M"
                ),
                booking["status"],
            ]
        )

    elements.append(
        Paragraph(
            "Upcoming Meetings",
            styles["Heading2"],
        )
    )

    meeting_table = Table(
        meeting_data,
        repeatRows=1,
    )

    meeting_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
            ]
        )
    )

    elements.append(meeting_table)

    document.build(elements)

    output.seek(0)

    return output